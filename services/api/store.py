from __future__ import annotations

import hashlib
import json
import mimetypes
import os
import re
import tempfile
from contextlib import contextmanager
from dataclasses import asdict, dataclass
from html.parser import HTMLParser
from pathlib import Path
from typing import Iterable

from filelock import FileLock
from docx import Document as DocxDocument
from psycopg.rows import dict_row
from psycopg.types.json import Jsonb
from pypdf import PdfReader

from packages.core.database import connect, database_dsn, vector_literal
from packages.core.knowledge_graph import query_terms
from packages.core.object_storage import ObjectStorage, object_storage_from_env


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".html", ".htm"}
TOKEN_PATTERN = re.compile(r"[0-9A-Za-z가-힣]+")


class _HTMLTextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        if data.strip():
            self.parts.append(data.strip())


@dataclass(frozen=True)
class Chunk:
    document_id: str
    filename: str
    chunk_index: int
    content: str
    page_number: int | None = None


@dataclass(frozen=True)
class DocumentRecord:
    document_id: str
    filename: str
    checksum: str
    size: int
    chunk_count: int
    version: int = 1


@dataclass(frozen=True)
class StoredDocument:
    document_id: str
    filename: str
    checksum: str
    size: int
    path: str
    object_key: str | None = None


def _normalize_text(text: str) -> str:
    # PostgreSQL text columns reject NUL bytes. Some PDF text extractors emit
    # them for embedded fonts or malformed content, so remove them before
    # chunking and before any value reaches full-text/vector indexing.
    return re.sub(r"\s+", " ", text.replace("\x00", "")).strip()


def extract_pages(path: Path) -> list[tuple[int | None, str]]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        pages = [(index, page.extract_text() or "") for index, page in enumerate(PdfReader(path).pages, 1)]
    elif suffix == ".docx":
        document = DocxDocument(path)
        pages = [(None, "\n".join(paragraph.text for paragraph in document.paragraphs))]
    elif suffix in {".html", ".htm"}:
        parser = _HTMLTextExtractor()
        parser.feed(path.read_text(encoding="utf-8", errors="replace"))
        pages = [(None, "\n".join(parser.parts))]
    elif suffix == ".txt":
        pages = [(None, path.read_text(encoding="utf-8", errors="replace"))]
    else:
        raise ValueError(f"지원하지 않는 파일 형식입니다: {suffix or '(확장자 없음)'}")

    normalized = [(page, _normalize_text(text)) for page, text in pages if _normalize_text(text)]
    if not normalized:
        if suffix == ".pdf":
            raise ValueError(
                "PDF에서 텍스트를 추출하지 못했습니다. "
                "텍스트 레이어가 없는 스캔 PDF일 수 있습니다. OCR 처리 후 다시 업로드하세요."
            )
        raise ValueError("문서에서 텍스트를 추출하지 못했습니다.")
    return normalized


def extract_text(path: Path) -> str:
    """하위 호환용 전체 본문 추출 함수."""
    return " ".join(text for _page, text in extract_pages(path))


def chunks_with_pages(path: Path) -> list[tuple[int | None, str]]:
    return [
        (page_number, chunk)
        for page_number, page_text in extract_pages(path)
        for chunk in split_text(page_text)
    ]


def split_text(text: str, chunk_size: int = 1000, overlap: int = 150) -> Iterable[str]:
    start = 0
    while start < len(text):
        end = min(start + chunk_size, len(text))
        if end < len(text):
            boundary = text.rfind(" ", start + chunk_size // 2, end)
            if boundary > start:
                end = boundary
        chunk = text[start:end].strip()
        if chunk:
            yield chunk
        if end >= len(text):
            break
        start = max(end - overlap, start + 1)


def _query_tokens(query: str) -> list[str]:
    return list(dict.fromkeys(token.lower() for token in TOKEN_PATTERN.findall(query) if len(token) > 1))


def _best_snippet(content: str, query: str, limit: int = 320) -> tuple[str, list[str]]:
    """질의 단어가 가장 많이 포함된 문장 주변을 짧은 근거 문장으로 반환한다."""
    terms = _query_tokens(query)
    sentences = [part.strip() for part in re.split(r"(?<=[.!?。])\s+|\n+", content) if part.strip()]
    if not sentences:
        sentences = [content.strip()]
    ranked = sorted(
        enumerate(sentences),
        key=lambda item: (sum(term in item[1].lower() for term in terms), -item[0]),
        reverse=True,
    )
    center = ranked[0][0]
    selected = " ".join(sentences[max(0, center - 1): center + 2]).strip()
    if len(selected) > limit:
        lowered = selected.lower()
        positions = [lowered.find(term) for term in terms if lowered.find(term) >= 0]
        start = max(0, (min(positions) if positions else 0) - limit // 3)
        selected = selected[start:start + limit].strip()
        if start:
            selected = "…" + selected
        if start + limit < len(content):
            selected += "…"
    matched = [term for term in terms if term in selected.lower()]
    return selected, matched


class DocumentStore:
    def __init__(self, data_dir: Path | None = None) -> None:
        configured = os.getenv("TIBERO_DOC_DATA_DIR")
        self.data_dir = data_dir or (Path(configured) if configured else Path.home() / ".tibero-doc" / "data")
        self.upload_dir = self.data_dir / "uploads"
        self.index_file = self.data_dir / "index.json"
        self.index_lock_file = self.data_dir / "index.lock"
        self.upload_dir.mkdir(parents=True, exist_ok=True)

    @contextmanager
    def _index_lock(self):
        with FileLock(str(self.index_lock_file)):
            yield

    def _load(self) -> dict:
        if not self.index_file.exists():
            return {"documents": [], "chunks": []}
        try:
            return json.loads(self.index_file.read_text(encoding="utf-8"))
        except (json.JSONDecodeError, OSError) as exc:
            raise RuntimeError(f"인덱스를 읽을 수 없습니다: {exc}") from exc

    def _save(self, index: dict) -> None:
        self.data_dir.mkdir(parents=True, exist_ok=True)
        descriptor, temporary_name = tempfile.mkstemp(dir=self.data_dir, prefix="index-", suffix=".tmp")
        try:
            with os.fdopen(descriptor, "w", encoding="utf-8") as temporary:
                json.dump(index, temporary, ensure_ascii=False, indent=2)
                temporary.flush()
                os.fsync(temporary.fileno())
            Path(temporary_name).replace(self.index_file)
        finally:
            Path(temporary_name).unlink(missing_ok=True)

    def save_upload(self, filename: str, content: bytes) -> StoredDocument:
        """Validate and persist an upload without performing expensive indexing."""
        safe_name = Path(filename).name
        suffix = Path(safe_name).suffix.lower()
        if suffix not in SUPPORTED_EXTENSIONS:
            raise ValueError(f"지원하지 않는 파일 형식입니다: {suffix or '(확장자 없음)'}")
        if not content:
            raise ValueError("빈 파일은 업로드할 수 없습니다.")

        checksum = hashlib.sha256(content).hexdigest()
        target = self.upload_dir / safe_name
        target.write_bytes(content)
        return StoredDocument(checksum[:16], safe_name, checksum, len(content), str(target))

    def ingest_bytes(self, filename: str, content: bytes) -> DocumentRecord:
        stored = self.save_upload(filename, content)
        target = Path(stored.path)
        try:
            return self._index_file(target)
        except Exception:
            target.unlink(missing_ok=True)
            raise

    def index_upload(self, filename: str) -> DocumentRecord:
        path = self.upload_dir / Path(filename).name
        if not path.is_file():
            raise ValueError(f"업로드 파일을 찾을 수 없습니다: {filename}")
        return self._index_file(path)

    def _index_file(self, path: Path) -> DocumentRecord:
        with self._index_lock():
            return self._index_file_unlocked(path)

    def _index_file_unlocked(self, path: Path) -> DocumentRecord:
        content = path.read_bytes()
        checksum = hashlib.sha256(content).hexdigest()
        document_id = checksum[:16]
        chunks = [
            Chunk(document_id, path.name, position, text, page_number)
            for position, (page_number, text) in enumerate(chunks_with_pages(path))
        ]
        record = DocumentRecord(document_id, path.name, checksum, len(content), len(chunks))

        index = self._load()
        index["documents"] = [item for item in index["documents"] if item["filename"] != path.name]
        index["chunks"] = [item for item in index["chunks"] if item["filename"] != path.name]
        index["documents"].append(asdict(record))
        index["chunks"].extend(asdict(chunk) for chunk in chunks)
        self._save(index)
        return record

    def search(
        self,
        query: str,
        top_k: int,
        query_vector: list[float] | None = None,
        model: str | None = None,
    ) -> list[dict]:
        terms = TOKEN_PATTERN.findall(query.lower())
        if not terms:
            return []

        results: list[dict] = []
        for chunk in self._load()["chunks"]:
            content = chunk["content"]
            lowered = content.lower()
            matches = sum(lowered.count(term) for term in terms)
            if matches == 0:
                continue
            coverage = sum(1 for term in set(terms) if term in lowered) / len(set(terms))
            density = matches / max(len(TOKEN_PATTERN.findall(lowered)), 1)
            score = min(1.0, coverage * 0.8 + density * 0.2)
            results.append({
                "document_id": chunk["document_id"],
                "filename": chunk["filename"],
                "chunk_index": chunk["chunk_index"],
                "content": content,
                "score": round(score, 6),
            })
        return sorted(results, key=lambda item: item["score"], reverse=True)[:top_k]

    def chunks_for_document(self, document_id: str) -> list[dict]:
        return [chunk for chunk in self._load()["chunks"] if chunk["document_id"] == document_id]

    def sync(self) -> dict[str, int]:
        with self._index_lock():
            return self._sync_unlocked()

    def _sync_unlocked(self) -> dict[str, int]:
        index = self._load()
        known = {item["filename"]: item for item in index["documents"]}
        current = {
            path.name: path
            for path in self.upload_dir.iterdir()
            if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS
        }
        added = updated = skipped = 0
        for filename, path in current.items():
            checksum = hashlib.sha256(path.read_bytes()).hexdigest()
            previous = known.get(filename)
            if previous is None:
                self._index_file_unlocked(path)
                added += 1
            elif previous["checksum"] != checksum:
                self._index_file_unlocked(path)
                updated += 1
            else:
                skipped += 1

        deleted_names = set(known) - set(current)
        if deleted_names:
            index = self._load()
            index["documents"] = [item for item in index["documents"] if item["filename"] not in deleted_names]
            index["chunks"] = [item for item in index["chunks"] if item["filename"] not in deleted_names]
            self._save(index)

        return {"added": added, "updated": updated, "deleted": len(deleted_names), "skipped": skipped}

    def stats(self) -> dict[str, int]:
        index = self._load()
        return {"documents": len(index["documents"]), "chunks": len(index["chunks"])}

    def list_documents(self, limit: int = 100, offset: int = 0) -> list[dict]:
        documents = sorted(self._load()["documents"], key=lambda item: item["filename"])
        return documents[offset:offset + limit]

    def get_document(self, document_id: str) -> dict | None:
        return next((item for item in self._load()["documents"] if item["document_id"] == document_id), None)

    def delete_document(self, document_id: str) -> bool:
        with self._index_lock():
            index = self._load()
            target = next((item for item in index["documents"] if item["document_id"] == document_id), None)
            if target is None:
                return False
            index["documents"] = [item for item in index["documents"] if item["document_id"] != document_id]
            index["chunks"] = [item for item in index["chunks"] if item["document_id"] != document_id]
            self._save(index)
            (self.upload_dir / target["filename"]).unlink(missing_ok=True)
            return True

    def versions(self, filename: str) -> list[dict]:
        record = next((item for item in self._load()["documents"] if item["filename"] == filename), None)
        return [record] if record else []


class OpenSQLDocumentStore(DocumentStore):
    """Document, chunk, metadata and hybrid-search repository backed by OpenSQL."""

    def __init__(self, dsn: str | None = None, data_dir: Path | None = None, workspace_id: str | None = None, user_id: str | None = None) -> None:
        super().__init__(data_dir)
        self.dsn = dsn or database_dsn()
        if not self.dsn:
            raise RuntimeError("TIBERO_DOC_DSN is not configured")
        self.workspace_id = workspace_id or "00000000-0000-0000-0000-000000000001"
        self.user_id = user_id or "00000000-0000-0000-0000-000000000001"
        self.objects: ObjectStorage = object_storage_from_env()

    def save_upload(self, filename: str, content: bytes) -> StoredDocument:
        stored = super().save_upload(filename, content)
        object_key = f"{self.workspace_id}/{stored.checksum}/{stored.filename}"
        self.objects.put(object_key, content, mimetypes.guess_type(stored.filename)[0] or "application/octet-stream")
        with connect(self.dsn) as connection:
            connection.execute(
                """INSERT INTO tibero_doc.object_lifecycle(object_key,document_id,storage_tier,last_accessed_at)
                   VALUES (%s,%s,'hot',now()) ON CONFLICT(object_key) DO UPDATE SET last_accessed_at=now()""",
                (object_key, stored.document_id),
            )
        return StoredDocument(stored.document_id, stored.filename, stored.checksum, stored.size, stored.path, object_key)

    def index_upload(self, filename: str, object_key: str | None = None) -> DocumentRecord:
        path = self.upload_dir / Path(filename).name
        if not path.is_file() and object_key:
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(self.objects.get(object_key))
        return super().index_upload(filename)

    def _index_file(self, path: Path) -> DocumentRecord:
        return self._index_file_unlocked(path)

    def _index_file_unlocked(self, path: Path) -> DocumentRecord:
        binary = path.read_bytes()
        checksum = hashlib.sha256(binary).hexdigest()
        document_id = hashlib.sha256(f"{self.workspace_id}:{checksum}".encode()).hexdigest()[:16]
        object_key = f"{self.workspace_id}/{checksum}/{path.name}"
        chunk_values = chunks_with_pages(path)
        with connect(self.dsn) as connection:
            previous = connection.execute(
                "SELECT document_id, checksum, size_bytes, chunk_count, version FROM tibero_doc.documents WHERE workspace_id = %s AND filename = %s",
                (self.workspace_id, path.name),
            ).fetchone()
            if previous and previous[1] == checksum:
                return DocumentRecord(previous[0], path.name, previous[1], previous[2], previous[3], previous[4])
            version = (previous[4] + 1) if previous else 1
            if previous:
                connection.execute("DELETE FROM tibero_doc.documents WHERE document_id = %s", (previous[0],))
            metadata = {"extension": path.suffix.lower(), "source": "upload"}
            connection.execute(
                """
                INSERT INTO tibero_doc.documents
                    (document_id, workspace_id, owner_user_id, filename, checksum, size_bytes, chunk_count, version, metadata, object_key)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                """,
                (document_id, self.workspace_id, self.user_id, path.name, checksum, len(binary), len(chunk_values), version, Jsonb(metadata), object_key),
            )
            with connection.cursor() as cursor:
                cursor.executemany(
                    "INSERT INTO tibero_doc.chunks (document_id, chunk_index, content, page_number) VALUES (%s, %s, %s, %s)",
                    [(document_id, index, content, page_number) for index, (page_number, content) in enumerate(chunk_values)],
                )
            connection.execute(
                """
                INSERT INTO tibero_doc.document_versions
                    (workspace_id, filename, version, document_id, checksum, size_bytes, chunk_count, metadata, object_key)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)
                ON CONFLICT (workspace_id, filename, version) DO NOTHING
                """,
                (self.workspace_id, path.name, version, document_id, checksum, len(binary), len(chunk_values), Jsonb(metadata), object_key),
            )
        return DocumentRecord(document_id, path.name, checksum, len(binary), len(chunk_values), version)

    def search(self, query: str, top_k: int, query_vector: list[float] | None = None, model: str | None = None) -> list[dict]:
        candidate_limit = max(top_k * 4, 20)
        keyword_patterns = [f"%{token}%" for token in _query_tokens(query)] or [f"%{query}%"]
        with connect(self.dsn) as connection:
            with connection.cursor(row_factory=dict_row) as cursor:
                cursor.execute(
                    """
                    SELECT c.document_id, d.filename, c.chunk_index, c.content, c.page_number,
                           ts_rank_cd(c.search_vector, websearch_to_tsquery('simple', %s)) AS rank
                      FROM tibero_doc.chunks c JOIN tibero_doc.documents d USING (document_id)
                     WHERE d.workspace_id = %s
                       AND (d.visibility='workspace' OR d.owner_user_id=%s OR EXISTS (
                         SELECT 1 FROM tibero_doc.document_acl a WHERE a.document_id=d.document_id
                           AND ((a.principal_type='user' AND a.principal_id=%s)
                             OR (a.principal_type='group' AND a.principal_id IN (SELECT group_id FROM tibero_doc.group_members WHERE user_id=%s)))))
                       AND (c.search_vector @@ websearch_to_tsquery('simple', %s) OR c.content ILIKE ANY(%s))
                     ORDER BY rank DESC, c.document_id, c.chunk_index LIMIT %s
                    """,
                    (query, self.workspace_id, self.user_id, self.user_id, self.user_id, query, keyword_patterns, candidate_limit),
                )
                keyword_rows = cursor.fetchall()
                vector_rows: list[dict] = []
                if query_vector and any(query_vector) and model:
                    literal = vector_literal(query_vector)
                    cursor.execute(
                        """
                        SELECT c.document_id, d.filename, c.chunk_index, c.content, c.page_number,
                               e.embedding <=> %s::vector AS distance
                          FROM tibero_doc.chunk_embeddings e
                          JOIN tibero_doc.chunks c USING (document_id, chunk_index)
                          JOIN tibero_doc.documents d USING (document_id)
                         WHERE d.workspace_id = %s
                           AND (d.visibility='workspace' OR d.owner_user_id=%s OR EXISTS (
                             SELECT 1 FROM tibero_doc.document_acl a WHERE a.document_id=d.document_id
                               AND ((a.principal_type='user' AND a.principal_id=%s)
                                 OR (a.principal_type='group' AND a.principal_id IN (SELECT group_id FROM tibero_doc.group_members WHERE user_id=%s)))))
                           AND e.model = %s AND e.dimensions = %s
                         ORDER BY e.embedding <=> %s::vector LIMIT %s
                        """,
                        (literal, self.workspace_id, self.user_id, self.user_id, self.user_id, model, len(query_vector), literal, candidate_limit),
                    )
                    vector_rows = cursor.fetchall()
                graph_rows: list[dict] = []
                terms = query_terms(query)
                if terms:
                    patterns = [f"%{term}%" for term in terms]
                    cursor.execute(
                        """
                        WITH matched AS (
                            SELECT entity_id FROM tibero_doc.entities
                             WHERE workspace_id=%s AND name ILIKE ANY(%s)
                        ), expanded AS (
                            SELECT entity_id FROM matched
                            UNION
                            SELECT CASE WHEN r.source_entity_id=m.entity_id
                                        THEN r.target_entity_id ELSE r.source_entity_id END
                              FROM matched m JOIN tibero_doc.relationships r
                                ON r.workspace_id=%s AND
                                   (r.source_entity_id=m.entity_id OR r.target_entity_id=m.entity_id)
                        )
                        SELECT c.document_id, d.filename, c.chunk_index, c.content, c.page_number,
                               array_agg(DISTINCT e.name) AS entities,
                               count(DISTINCT e.entity_id) AS entity_matches
                          FROM expanded x
                          JOIN tibero_doc.entities e ON e.entity_id=x.entity_id
                          JOIN tibero_doc.document_entities de ON de.entity_id=x.entity_id
                          JOIN tibero_doc.chunks c ON c.document_id=de.document_id AND c.chunk_index=de.chunk_index
                          JOIN tibero_doc.documents d ON d.document_id=c.document_id
                         WHERE d.workspace_id=%s
                           AND (d.visibility='workspace' OR d.owner_user_id=%s OR EXISTS (
                             SELECT 1 FROM tibero_doc.document_acl a WHERE a.document_id=d.document_id
                               AND ((a.principal_type='user' AND a.principal_id=%s)
                                 OR (a.principal_type='group' AND a.principal_id IN
                                     (SELECT group_id FROM tibero_doc.group_members WHERE user_id=%s)))))
                         GROUP BY c.document_id, d.filename, c.chunk_index, c.content, c.page_number
                         ORDER BY entity_matches DESC, c.document_id, c.chunk_index LIMIT %s
                        """,
                        (self.workspace_id, patterns, self.workspace_id, self.workspace_id,
                         self.user_id, self.user_id, self.user_id, candidate_limit),
                    )
                    graph_rows = cursor.fetchall()
        # 해시 임베딩은 파이프라인 데모용이며 의미 유사도로 사용하지 않는다.
        if model and model.startswith("local-hash-"):
            vector_rows = []
        minimum_similarity = float(os.getenv("VECTOR_MIN_SIMILARITY", "0.50"))
        vector_rows = [row for row in vector_rows if 1.0 - float(row["distance"]) >= minimum_similarity]

        combined: dict[tuple[str, int], dict] = {}
        for rank, row in enumerate(keyword_rows, start=1):
            key = (row["document_id"], row["chunk_index"])
            combined[key] = {
                "document_id": row["document_id"], "filename": row["filename"],
                "chunk_index": row["chunk_index"], "content": row["content"],
                "page_number": row.get("page_number"),
                "keyword_rank": rank, "vector_rank": None, "_rrf": 1 / (60 + rank),
                "graph_rank": None, "entities": [], "vector_similarity": None,
            }
        for rank, row in enumerate(vector_rows, start=1):
            key = (row["document_id"], row["chunk_index"])
            item = combined.setdefault(key, {
                "document_id": row["document_id"], "filename": row["filename"],
                "chunk_index": row["chunk_index"], "content": row["content"],
                "page_number": row.get("page_number"),
                "keyword_rank": None, "vector_rank": None, "_rrf": 0.0,
                "graph_rank": None, "entities": [], "vector_similarity": None,
            })
            item["vector_rank"] = rank
            item["vector_similarity"] = round(1.0 - float(row["distance"]), 6)
            item["_rrf"] += 1 / (60 + rank)
        for rank, row in enumerate(graph_rows, start=1):
            key = (row["document_id"], row["chunk_index"])
            item = combined.setdefault(key, {
                "document_id": row["document_id"], "filename": row["filename"],
                "chunk_index": row["chunk_index"], "content": row["content"],
                "page_number": row.get("page_number"),
                "keyword_rank": None, "vector_rank": None, "graph_rank": None,
                "entities": [], "vector_similarity": None, "_rrf": 0.0,
            })
            item["graph_rank"] = rank
            item["entities"] = list(row["entities"] or [])
            item["_rrf"] += 1 / (60 + rank)
        # 같은 문서의 청크를 하나의 결과 카드로 묶고 상위 근거만 보존한다.
        grouped: dict[str, dict] = {}
        ordered_chunks = sorted(combined.values(), key=lambda item: item["_rrf"], reverse=True)
        for item in ordered_chunks:
            snippet, matched_terms = _best_snippet(item["content"], query)
            passage = {
                "chunk_index": item["chunk_index"], "page_number": item.get("page_number"),
                "snippet": snippet, "matched_terms": matched_terms,
                "keyword_rank": item.get("keyword_rank"), "vector_rank": item.get("vector_rank"),
                "graph_rank": item.get("graph_rank"), "vector_similarity": item.get("vector_similarity"),
            }
            document = grouped.setdefault(item["document_id"], {
                "document_id": item["document_id"], "filename": item["filename"],
                "chunk_index": item["chunk_index"], "page_number": item.get("page_number"),
                "content": snippet, "snippet": snippet, "matched_terms": matched_terms,
                "keyword_rank": item.get("keyword_rank"), "vector_rank": item.get("vector_rank"),
                "graph_rank": item.get("graph_rank"), "vector_similarity": item.get("vector_similarity"),
                "entities": list(item.get("entities", [])), "passages": [], "_rrf": 0.0,
            })
            document["_rrf"] += item["_rrf"]
            document["entities"] = list(dict.fromkeys(document["entities"] + item.get("entities", [])))
            for rank_name in ("keyword_rank", "vector_rank", "graph_rank"):
                rank_value = item.get(rank_name)
                if rank_value is not None:
                    current_rank = document.get(rank_name)
                    document[rank_name] = min(current_rank, rank_value) if current_rank is not None else rank_value
            if len(document["passages"]) < int(os.getenv("SEARCH_PASSAGES_PER_DOCUMENT", "3")):
                document["passages"].append(passage)

        results = sorted(grouped.values(), key=lambda item: item["_rrf"], reverse=True)[:top_k]
        for item in results:
            item.pop("_rrf", None)
            signals = [p["vector_similarity"] for p in item["passages"] if p.get("vector_similarity") is not None]
            lexical = max((len(p["matched_terms"]) / max(len(_query_tokens(query)), 1) for p in item["passages"]), default=0.0)
            item["relevance"] = round(max([lexical, *signals]), 6)
            item["relevance_label"] = "매우 높음" if item["relevance"] >= 0.75 else "높음" if item["relevance"] >= 0.55 else "보통"
            item["score"] = item["relevance"]
        minimum_relevance = float(os.getenv("SEARCH_MIN_RELEVANCE", "0.35"))
        return [
            item for item in results
            if item["relevance"] >= minimum_relevance or item.get("graph_rank") is not None
        ]

    def chunks_for_document(self, document_id: str) -> list[dict]:
        with connect(self.dsn) as connection:
            rows = connection.execute(
                "SELECT document_id, chunk_index, content, page_number FROM tibero_doc.chunks WHERE document_id = %s ORDER BY chunk_index",
                (document_id,),
            ).fetchall()
        return [{"document_id": row[0], "chunk_index": row[1], "content": row[2], "page_number": row[3]} for row in rows]

    def list_documents(self, limit: int = 100, offset: int = 0) -> list[dict]:
        with connect(self.dsn) as connection:
            with connection.cursor(row_factory=dict_row) as cursor:
                cursor.execute(
                    """
                    SELECT document_id, filename, checksum, size_bytes AS size, chunk_count,
                           version, metadata, object_key, created_at, updated_at
                      FROM tibero_doc.documents d WHERE workspace_id = %s
                       AND (d.visibility='workspace' OR d.owner_user_id=%s OR EXISTS (
                         SELECT 1 FROM tibero_doc.document_acl a WHERE a.document_id=d.document_id
                           AND ((a.principal_type='user' AND a.principal_id=%s)
                             OR (a.principal_type='group' AND a.principal_id IN (SELECT group_id FROM tibero_doc.group_members WHERE user_id=%s)))))
                      ORDER BY updated_at DESC LIMIT %s OFFSET %s
                    """,
                    (self.workspace_id, self.user_id, self.user_id, self.user_id, limit, offset),
                )
                return [dict(row) for row in cursor.fetchall()]

    def get_document(self, document_id: str) -> dict | None:
        with connect(self.dsn) as connection:
            with connection.cursor(row_factory=dict_row) as cursor:
                cursor.execute(
                    """
                    SELECT document_id, filename, checksum, size_bytes AS size, chunk_count,
                           version, metadata, object_key, created_at, updated_at
                      FROM tibero_doc.documents WHERE workspace_id = %s AND document_id = %s
                    """,
                    (self.workspace_id, document_id),
                )
                row = cursor.fetchone()
                return dict(row) if row else None

    def delete_document(self, document_id: str) -> bool:
        with connect(self.dsn) as connection:
            row = connection.execute(
                "DELETE FROM tibero_doc.documents WHERE workspace_id = %s AND document_id = %s RETURNING filename",
                (self.workspace_id, document_id),
            ).fetchone()
        if row:
            (self.upload_dir / row[0]).unlink(missing_ok=True)
        return row is not None

    def versions(self, filename: str) -> list[dict]:
        with connect(self.dsn) as connection:
            with connection.cursor(row_factory=dict_row) as cursor:
                cursor.execute(
                    """
                    SELECT filename, version, document_id, checksum, size_bytes AS size,
                           chunk_count, metadata, created_at
                      FROM tibero_doc.document_versions WHERE workspace_id = %s AND filename = %s ORDER BY version DESC
                    """,
                    (self.workspace_id, filename),
                )
                return [dict(row) for row in cursor.fetchall()]

    def download(self, document_id: str) -> tuple[bytes | None, str | None, str]:
        document = self.get_document(document_id)
        if not document or not document.get("object_key"):
            raise FileNotFoundError(document_id)
        url = self.objects.download_url(document["object_key"])
        return (None, url, document["filename"]) if url else (self.objects.get(document["object_key"]), None, document["filename"])

    def _sync_unlocked(self) -> dict:
        known = {item["filename"]: item for item in self.list_documents(10000, 0)}
        current = {path.name: path for path in self.upload_dir.iterdir() if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS}
        added = updated = skipped = 0
        changed_ids: list[str] = []
        for filename, path in current.items():
            checksum = hashlib.sha256(path.read_bytes()).hexdigest()
            previous = known.get(filename)
            if previous is None or previous["checksum"] != checksum:
                record = self._index_file_unlocked(path)
                changed_ids.append(record.document_id)
                if previous is None:
                    added += 1
                else:
                    updated += 1
            else:
                skipped += 1
        deleted_names = set(known) - set(current)
        for filename in deleted_names:
            self.delete_document(known[filename]["document_id"])
        return {"added": added, "updated": updated, "deleted": len(deleted_names), "skipped": skipped, "document_ids": changed_ids}

    def stats(self) -> dict:
        with connect(self.dsn) as connection:
            row = connection.execute(
                """
                SELECT (SELECT count(*) FROM tibero_doc.documents WHERE workspace_id = %s),
                       (SELECT count(*) FROM tibero_doc.chunks c JOIN tibero_doc.documents d USING(document_id) WHERE d.workspace_id = %s),
                       (SELECT count(*) FROM tibero_doc.chunk_embeddings e JOIN tibero_doc.documents d USING(document_id) WHERE d.workspace_id = %s),
                       (SELECT count(*) FROM tibero_doc.outbox_events WHERE published_at IS NULL)
                """,
                (self.workspace_id, self.workspace_id, self.workspace_id),
            ).fetchone()
        return {"documents": row[0], "chunks": row[1], "embeddings": row[2], "pending_events": row[3], "storage": "opensql"}


def document_store_from_env(workspace_id: str | None = None, user_id: str | None = None) -> DocumentStore:
    return OpenSQLDocumentStore(workspace_id=workspace_id, user_id=user_id) if database_dsn() else DocumentStore()
